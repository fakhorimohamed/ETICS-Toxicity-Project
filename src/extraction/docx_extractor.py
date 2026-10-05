# extractors/docx_extractor.py
from pathlib import Path
from docx import Document as DocxDocument
from src.extraction.DocumentBase import Document


def extract_docx(file_path: str) -> Document:
    """Extrait le texte d'un fichier DOCX avec métadonnées."""
    path = Path(file_path)
    
    doc = DocxDocument(file_path)
    
    # Extraction du texte des paragraphes
    paragraphs = []
    for para in doc.paragraphs:
        if para.text.strip():
            paragraphs.append(para.text)
    
    full_text = "\n\n".join(paragraphs)
    
    # Extraction des tableaux (souvent présents dans les contrats)
    tables_text = []
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                tables_text.append(row_text)
    
    if tables_text:
        full_text += "\n\n--- TABLEAUX ---\n\n" + "\n".join(tables_text)
    
    # Métadonnées du DOCX
    core_props = doc.core_properties
    
    return Document(
        text=full_text.strip(),
        source_filename=path.name,
        source_path=str(path.absolute()),
        file_format="docx",
        file_size_bytes=path.stat().st_size,
        number_of_pages=None,  # DOCX n'a pas de nombre de pages natif
        metadata={
            "docx_title": core_props.title or "",
            "docx_author": core_props.author or "",
            "docx_subject": core_props.subject or "",
            "docx_keywords": core_props.keywords or "",
            "docx_created": str(core_props.created) if core_props.created else "",
            "docx_modified": str(core_props.modified) if core_props.modified else "",
            "docx_last_modified_by": core_props.last_modified_by or "",
            "docx_paragraph_count": len(doc.paragraphs),
            "docx_table_count": len(doc.tables),
        }
    )